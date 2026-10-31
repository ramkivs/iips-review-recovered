#!/usr/bin/env python3
"""Verify the reassessment bundle locally. Requires openpyxl.

    python verify_reassessment.py

Checks shape, headers, numbering, CSV/XLSX consistency, pending traceability,
evidence IDs, filters/tables, and SHA256SUMS if present. No network required.
"""
from __future__ import annotations
import csv, hashlib, json, pathlib, sys
from openpyxl import load_workbook
import baseline_matrix_data as base

HERE = pathlib.Path(__file__).resolve().parent
fails=[]
def check(name, ok, detail=''):
    print(('PASS' if ok else 'FAIL')+': '+name+((' — '+detail) if detail else ''))
    if not ok: fails.append(name)

xlsx=HERE/'IIPS_V3_CAPABILITY_MATRIX_51x17_REASSESSED_2026-10-31.xlsx'
mcsv=HERE/'IIPS_V3_CAPABILITY_MATRIX_51x17_REASSESSED_2026-10-31.csv'
pcsv=HERE/'IIPS_V3_OUTSTANDING_PENDING_2026-10-31.csv'
wb=load_workbook(xlsx)
ws=wb['Updated Capability Matrix']; wp=wb['Outstanding-Pending Items']
check('workbook sheets', wb.sheetnames[:2]==['Updated Capability Matrix','Outstanding-Pending Items'], ', '.join(wb.sheetnames))
check('matrix shape 51x17', ws.max_row==52 and ws.max_column==17, f'{ws.max_row-1}x{ws.max_column}')
headers=[ws.cell(1,c).value for c in range(1,18)]
check('headers preserved', headers==base.HEADERS)
rows=[[ws.cell(r,c).value for c in range(1,18)] for r in range(2,53)]
check('row numbering', [r[0] for r in rows]==list(range(1,52)))
check('no empty matrix cells', all(str(c).strip() for r in rows for c in r))
with mcsv.open(newline='',encoding='utf-8') as f: mr=list(csv.reader(f))
check('matrix CSV shape', len(mr)==52 and all(len(r)==17 for r in mr))
check('matrix CSV/XLSX equal', mr[0]==headers and [[str(c) for c in r] for r in rows]==mr[1:])
with pcsv.open(newline='',encoding='utf-8') as f: pr=list(csv.reader(f))
check('pending CSV shape', len(pr)==wp.max_row and all(len(r)==16 for r in pr), f'{len(pr)-1} items x 16')
pids=[r[0] for r in pr[1:]]
check('pending IDs unique', len(pids)==len(set(pids)))
anchors=' '.join(str(r[16]) for r in rows)
missing=[i for i in pids if i not in anchors]
check('pending traceability', not missing, ','.join(missing))
check('matrix filter', ws.auto_filter.ref=='A1:Q52')
check('pending filter', wp.auto_filter.ref==f'A1:P{wp.max_row}')
check('matrix table', 'CapabilityMatrix' in ws.tables)
check('pending table', 'PendingItems' in wp.tables)
er=wb['Evidence Register']; eids=[er.cell(r,1).value for r in range(2,er.max_row+1)]
check('evidence IDs unique', len(eids)==len(set(eids)))
# every evidence ID cited in matrix/pending/report exists
cited=set()
import re
text=anchors+' '+' '.join(' '.join(str(c) for c in r) for r in pr[1:])+(HERE/'RECONCILIATION_REPORT.md').read_text()
for m in re.findall(r'E\d{3}', text): cited.add(m)
check('cited evidence IDs exist', cited.issubset(set(eids)), ','.join(sorted(cited-set(eids))))
sums=HERE/'SHA256SUMS.txt'
if sums.exists():
    ok=True; bad=[]
    for line in sums.read_text().splitlines():
        if not line.strip(): continue
        h,rel=line.split('  ',1); p=HERE/rel
        if not p.exists() or hashlib.sha256(p.read_bytes()).hexdigest()!=h: ok=False; bad.append(rel)
    check('SHA256SUMS', ok, ','.join(bad))
check('overall', not fails)
sys.exit(1 if fails else 0)
